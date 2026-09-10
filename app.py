from __future__ import annotations

import calendar
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Gestión de Recaudo y Asignación", page_icon="📊", layout="wide")

MESES = {1:"Enero",2:"Febrero",3:"Marzo",4:"Abril",5:"Mayo",6:"Junio",7:"Julio",8:"Agosto",9:"Septiembre",10:"Octubre",11:"Noviembre",12:"Diciembre"}
MESES_ABR = {"ENE":1,"FEB":2,"MAR":3,"ABR":4,"MAY":5,"JUN":6,"JUL":7,"AGO":8,"SEP":9,"OCT":10,"NOV":11,"DIC":12,
             "JAN":1,"APR":4,"AUG":8,"DEC":12}

ALIASES = {
    "obligacion": ["obligacion", "obligacion ics", "numero obligacion", "nro obligacion", "credito"],
    "cedula": ["cedula sin digito", "ced sin dig", "cedula", "documento", "identificacion"],
    "cliente": ["cliente", "nombre cliente", "nombre"],
    "tipo": ["estado oblig", "tipo cartera", "cartera", "marca castigo", "filtro vencida"],
    "recaudo": ["sum of valor consignacion", "valor consignacion", "valor recaudo", "valor pago", "recaudo"],
    "fecha_pago": ["fecha recaudo", "fecha de pago", "fecha pago"],
    "saldo": ["saldo oblig", "saldo pendiente", "saldo de capital", "capital total", "capital cliente"],
    "gestor": ["analista negociacion", "gestor", "gestor actual", "analista"],
}


def init_state() -> None:
    defaults = {
        "config_guardada": False, "anio": date.today().year, "mes": date.today().month,
        "meta_castigada": 0.0, "meta_vencida": 0.0,
        "gestores": pd.DataFrame(columns=["Gestor", "Prioridad", "Peso (%)", "Activo"]),
        "recaudo_limpio": None, "pagos_obligaciones": set(), "pagos_cedulas": set(),
        "resumen_metas": None, "fecha_corte": None, "calidad_recaudo": [],
        "asignacion_limpia": None, "resultado_asignacion": None,
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


def strip_accents(value: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", str(value)) if not unicodedata.combining(c))


def canon(value: str) -> str:
    s = strip_accents(value).lower().strip()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def normalizar_id(serie: pd.Series) -> pd.Series:
    def one(x):
        if pd.isna(x): return ""
        s = str(x).strip()
        s = re.sub(r"\.0$", "", s)
        s = re.sub(r"\s+", "", s)
        return s
    return serie.map(one).astype("string")


def limpiar_monetario(serie: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(serie):
        return pd.to_numeric(serie, errors="coerce")
    def one(x):
        if pd.isna(x): return None
        s = str(x).strip().replace("$", "").replace(" ", "")
        if not s: return None
        # Formatos colombianos y estándar.
        if "," in s and "." in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "," in s:
            tail = s.split(",")[-1]
            s = s.replace(",", ".") if len(tail) <= 2 else s.replace(",", "")
        return s
    return pd.to_numeric(serie.map(one), errors="coerce")


def clasificar_cartera(serie: pd.Series) -> pd.Series:
    def one(x):
        s = canon(x)
        if any(t in s for t in ["cast", "castig"]): return "Castigada"
        if any(t in s for t in ["venc", "mora"]): return "Vencida"
        return "Sin clasificar"
    return serie.fillna("").map(one)


def detectar_columna(columns, key: str):
    norm = {c: canon(c) for c in columns}
    aliases = [canon(x) for x in ALIASES[key]]
    for c, n in norm.items():
        if n in aliases: return c
    for c, n in norm.items():
        if any(a in n or n in a for a in aliases if len(a) >= 5): return c
    return None


def fecha_desde_nombre(nombre: str):
    stem = Path(nombre).stem.upper()
    m = re.search(r"(?<!\d)(\d{2})(\d{2})(\d{4})(?!\d)", stem)
    if m:
        try: return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError: pass
    m = re.search(r"(?<!\d)(\d{2})([A-Z]{3})(\d{4})(?!\d)", stem)
    if m and m.group(2) in MESES_ABR:
        try: return date(int(m.group(3)), MESES_ABR[m.group(2)], int(m.group(1)))
        except ValueError: pass
    m = re.search(r"(?<!\d)(20\d{2})[-_]?([01]\d)(?:[-_]?([0-3]\d))?(?!\d)", stem)
    if m:
        try: return date(int(m.group(1)), int(m.group(2)), int(m.group(3) or 1))
        except ValueError: pass
    return None


@st.cache_data(show_spinner=False)
def hojas_excel(raw: bytes):
    return pd.ExcelFile(BytesIO(raw)).sheet_names


@st.cache_data(show_spinner=False)
def leer_excel(raw: bytes, hoja: str, header=0):
    return pd.read_excel(BytesIO(raw), sheet_name=hoja, header=header)


def selector_columna(label, columns, detectada=None, required=True, key=None):
    opciones = list(columns)
    if not required: opciones = ["(No disponible)"] + opciones
    idx = opciones.index(detectada) if detectada in opciones else 0
    return st.selectbox(label, opciones, index=idx, key=key)


def pesos_automaticos(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    activos = out["Activo"].fillna(False).astype(bool)
    if activos.sum() == 0: return out
    pr = pd.to_numeric(out.loc[activos, "Prioridad"], errors="coerce")
    if pr.isna().any(): return out
    n = len(pr)
    bruto = n - pr + 1
    bruto = bruto.clip(lower=0.01)
    out.loc[activos, "Peso (%)"] = (bruto / bruto.sum() * 100).round(4).values
    out.loc[~activos, "Peso (%)"] = 0.0
    return out


def validar_gestores(df: pd.DataFrame):
    errores=[]
    if df.empty: return ["Registre al menos un gestor."]
    d=df.copy(); d["Gestor"]=d["Gestor"].fillna("").astype(str).str.strip()
    activos=d[d["Activo"].fillna(False).astype(bool)].copy()
    if activos.empty: errores.append("Debe existir al menos un gestor activo.")
    if d.loc[d["Gestor"]!="", "Gestor"].str.casefold().duplicated().any(): errores.append("Existen gestores duplicados.")
    p=pd.to_numeric(activos["Prioridad"], errors="coerce")
    w=pd.to_numeric(activos["Peso (%)"], errors="coerce")
    if p.isna().any() or (p<=0).any(): errores.append("Las prioridades activas deben ser enteros positivos.")
    if p.duplicated().any(): errores.append("Las posiciones de prioridad activas no pueden repetirse.")
    if w.isna().any() or (w<=0).any(): errores.append("Los pesos de gestores activos deben ser mayores que cero.")
    if not w.empty and abs(w.sum()-100)>0.05: errores.append(f"Los pesos de gestores activos suman {w.sum():.2f}% y deben sumar 100%.")
    return errores


def procesar_recaudo(df, mapeo, anio, mes):
    x=pd.DataFrame({
        "Obligacion": normalizar_id(df[mapeo["obligacion"]]),
        "Cedula": normalizar_id(df[mapeo["cedula"]]),
        "Tipo_original": df[mapeo["tipo"]].astype("string"),
        "Recaudo": limpiar_monetario(df[mapeo["recaudo"]]),
        "Fecha_pago": pd.to_datetime(df[mapeo["fecha_pago"]], errors="coerce", dayfirst=True),
    })
    x["Tipo_cartera"]=clasificar_cartera(x["Tipo_original"])
    calidad=[]
    calidad.append(("Obligación vacía", int((x.Obligacion=="").sum()), "Registros sin número de obligación"))
    calidad.append(("Cédula vacía", int((x.Cedula=="").sum()), "Registros sin cédula"))
    calidad.append(("Recaudo no numérico", int(x.Recaudo.isna().sum()), "Valores que no pudieron convertirse"))
    calidad.append(("Fecha inválida", int(x.Fecha_pago.isna().sum()), "Fechas que no pudieron convertirse"))
    calidad.append(("Cartera sin clasificar", int((x.Tipo_cartera=="Sin clasificar").sum()), "Valores distintos de Castigada/Vencida"))
    valid=x[(x.Fecha_pago.dt.year==anio)&(x.Fecha_pago.dt.month==mes)&x.Recaudo.notna()&(x.Recaudo>0)&(x.Obligacion!="")].copy()
    # Elimina filas idénticas. Si existen múltiples movimientos legítimos por obligación/fecha/valor, conserva uno por combinación exacta.
    dup=valid.duplicated(subset=["Obligacion","Cedula","Tipo_cartera","Recaudo","Fecha_pago"], keep="first")
    calidad.append(("Duplicados exactos de recaudo", int(dup.sum()), "Se excluyen antes de sumar"))
    valid=valid[~dup].copy()
    return valid, calidad


def resumen_recaudo(x, meta_cast, meta_venc, corte):
    rows=[]
    for tipo, meta in [("Castigada",meta_cast),("Vencida",meta_venc)]:
        d=x[x.Tipo_cartera==tipo]
        rec=float(d.Recaudo.sum()); meta=float(meta)
        rows.append({"Tipo de cartera":tipo,"Meta":meta,"Recaudo acumulado":rec,"Brecha pendiente":max(meta-rec,0),
                     "Cumplimiento":rec/meta if meta>0 else 0,"Obligaciones con pago":d.Obligacion.nunique(),"Clientes con pago":d.Cedula.replace("",pd.NA).nunique()})
    return pd.DataFrame(rows)


def procesar_asignacion(df, mapeo):
    cliente = df[mapeo["cliente"]].astype("string") if mapeo.get("cliente") and mapeo["cliente"]!="(No disponible)" else pd.Series("", index=df.index, dtype="string")
    gestor = df[mapeo["gestor"]].astype("string") if mapeo.get("gestor") and mapeo["gestor"]!="(No disponible)" else pd.Series("", index=df.index, dtype="string")
    x=pd.DataFrame({"Obligacion":normalizar_id(df[mapeo["obligacion"]]),"Cedula":normalizar_id(df[mapeo["cedula"]]),"Cliente":cliente,
                    "Tipo_original":df[mapeo["tipo"]].astype("string"),"Saldo":limpiar_monetario(df[mapeo["saldo"]]),"Gestor_actual":gestor})
    x["Tipo_cartera"]=clasificar_cartera(x.Tipo_original)
    calidad=[]
    calidad.append(("Obligación vacía",int((x.Obligacion=="").sum()),"Requiere revisión"))
    calidad.append(("Cédula vacía",int((x.Cedula=="").sum()),"Requiere revisión"))
    calidad.append(("Saldo no numérico",int(x.Saldo.isna().sum()),"Requiere revisión"))
    calidad.append(("Cartera sin clasificar",int((x.Tipo_cartera=="Sin clasificar").sum()),"Requiere revisión"))
    dups=x[(x.Obligacion!="")].duplicated("Obligacion", keep=False)
    calidad.append(("Obligaciones duplicadas",int(dups.sum()),"Se conserva una sola fila por obligación, priorizando el mayor saldo"))
    # Una obligación nunca puede aparecer dos veces en la salida.
    x=x.sort_values("Saldo",ascending=False).drop_duplicates("Obligacion",keep="first")
    return x, calidad


def priorizar_obligaciones(base, pagos_obl, pagos_ced, resumen_metas, admitir_parciales=False):
    x=base.copy()
    x["Tiene_pago_obligacion"]=x.Obligacion.isin(pagos_obl)
    # Regla FINESA solicitada: si una obligación tiene pago, se excluyen todas las obligaciones de la misma cédula.
    x["Cedula_con_pago"]=x.Cedula.isin(pagos_ced)
    x["Estado_pago_mes"]=x["Cedula_con_pago"].map({True:"Cédula con pago en el mes",False:"Sin pago en el mes"})
    x["Requiere_revision"]=(x.Obligacion=="")|(x.Cedula=="")|x.Saldo.isna()|(x.Saldo<=0)|(x.Tipo_cartera=="Sin clasificar")
    cumplimiento={r["Tipo de cartera"]:r["Cumplimiento"] for _,r in resumen_metas.iterrows()} if resumen_metas is not None else {}
    x["Cumplimiento_cartera"]=x.Tipo_cartera.map(cumplimiento).fillna(1.0)
    if admitir_parciales:
        elegible=~x.Requiere_revision
    else:
        elegible=(~x.Requiere_revision)&(~x.Cedula_con_pago)
    e=x[elegible].copy()
    e["Prioridad_pago"]=e.Cedula_con_pago.astype(int)  # 0 primero
    e=e.sort_values(["Prioridad_pago","Cumplimiento_cartera","Saldo"],ascending=[True,True,False]).reset_index(drop=True)
    e["Posicion_prioridad"]=range(1,len(e)+1)
    e["Motivo_priorizacion"]=e.apply(lambda r: ("Sin pago en el mes; mayor saldo y cartera con menor cumplimiento" if not r.Cedula_con_pago else "Pago parcial admitido; prioridad posterior"),axis=1)
    excl=x[~elegible].copy()
    def motivo(r):
        if r.Requiere_revision: return "Inconsistencia de datos / requiere revisión"
        if r.Cedula_con_pago: return "Excluida: la cédula registra al menos una obligación con pago en el mes"
        return "No elegible"
    excl["Motivo_exclusion"]=excl.apply(motivo,axis=1)
    return e, excl


def aplicar_alcance(e, modo, valor=None):
    if e.empty: return e.copy()
    if modo=="Todas las obligaciones elegibles": return e.copy()
    if modo=="Cantidad máxima de obligaciones": return e.head(int(valor)).copy()
    if modo=="Hasta alcanzar un monto total de cartera":
        limite=float(valor); c=e.Saldo.cumsum(); n=max(1,int((c<=limite).sum()))
        if n < len(e) and c.iloc[n-1] < limite: n+=1
        return e.head(n).copy()
    return e.copy()


def distribuir_obligaciones(e, gestores, max_por_gestor=None):
    if e.empty: return e.copy(), pd.DataFrame()
    g=gestores[gestores.Activo.fillna(False).astype(bool)].copy()
    g["Prioridad"]=pd.to_numeric(g.Prioridad); g["Peso (%)"]=pd.to_numeric(g["Peso (%)"])
    g=g.sort_values("Prioridad").reset_index(drop=True)
    total=float(e.Saldo.sum()); g["Monto objetivo"]=total*g["Peso (%)"]/100; g["Monto asignado"]=0.0; g["Cantidad obligaciones"]=0
    out=e.sort_values("Saldo",ascending=False).copy(); asign=[]
    # Primera ronda: las obligaciones más altas siguen el orden de prioridad de gestores.
    for idx, (_, row) in enumerate(out.iterrows()):
        elegibles=g.index if not max_por_gestor else g.index[g["Cantidad obligaciones"] < int(max_por_gestor)]
        if len(elegibles)==0: asign.append(None); continue
        if idx < len(g):
            pref=idx
            gi=pref if pref in elegibles else elegibles[0]
        else:
            cand=g.loc[elegibles].copy()
            cand["Deficit relativo"]=(cand["Monto objetivo"]-cand["Monto asignado"])/cand["Monto objetivo"].replace(0,pd.NA)
            cand["Deficit relativo"]=cand["Deficit relativo"].fillna(-999)
            gi=cand.sort_values(["Deficit relativo","Prioridad"],ascending=[False,True]).index[0]
        asign.append(g.at[gi,"Gestor"]); g.at[gi,"Monto asignado"]+=float(row.Saldo); g.at[gi,"Cantidad obligaciones"]+=1
    out["Gestor_asignado"]=asign
    prio=dict(zip(g.Gestor,g.Prioridad)); out["Prioridad_gestor"]=out.Gestor_asignado.map(prio)
    g["Porcentaje real"]=g["Monto asignado"]/total if total else 0
    g["Desviacion"]=g["Monto asignado"]-g["Monto objetivo"]
    g["Saldo promedio asignado"]=g["Monto asignado"]/g["Cantidad obligaciones"].replace(0,pd.NA)
    return out[out.Gestor_asignado.notna()].copy(), g


def formato_pesos(v): return f"${v:,.0f}".replace(",", ".")


def crear_excel(asig, resumen_g, resumen_m, exclusiones, calidad, fecha_asig):
    bio=BytesIO()
    det=pd.DataFrame({
        "Número de obligación":asig.get("Obligacion",pd.Series(dtype=str)),"Cédula":asig.get("Cedula",pd.Series(dtype=str)),"Nombre del cliente":asig.get("Cliente",pd.Series(dtype=str)),
        "Tipo de cartera":asig.get("Tipo_cartera",pd.Series(dtype=str)),"Saldo pendiente":asig.get("Saldo",pd.Series(dtype=float)),"Estado de pago del mes":asig.get("Estado_pago_mes",pd.Series(dtype=str)),
        "Posición de prioridad":asig.get("Posicion_prioridad",pd.Series(dtype=int)),"Motivo de priorización":asig.get("Motivo_priorizacion",pd.Series(dtype=str)),"Gestor asignado":asig.get("Gestor_asignado",pd.Series(dtype=str)),
        "Prioridad del gestor":asig.get("Prioridad_gestor",pd.Series(dtype=float)),"Fecha de asignación":fecha_asig,
    })
    rg=resumen_g.rename(columns={"Peso (%)":"Porcentaje objetivo"})[["Gestor","Prioridad","Porcentaje objetivo","Monto objetivo","Monto asignado","Porcentaje real","Desviacion","Cantidad obligaciones","Saldo promedio asignado"]] if not resumen_g.empty else pd.DataFrame(columns=["Gestor","Prioridad","Porcentaje objetivo","Monto objetivo","Monto asignado","Porcentaje real","Desviacion","Cantidad obligaciones","Saldo promedio asignado"])
    rm=resumen_m.copy() if resumen_m is not None else pd.DataFrame(columns=["Tipo de cartera","Meta","Recaudo acumulado","Brecha pendiente","Cumplimiento"])
    if "Cumplimiento" in rm: rm=rm.rename(columns={"Cumplimiento":"Porcentaje de cumplimiento"})
    rm["Fecha de corte"]=st.session_state.fecha_corte
    ex=pd.DataFrame({"Número de obligación":exclusiones.get("Obligacion",pd.Series(dtype=str)),"Cédula":exclusiones.get("Cedula",pd.Series(dtype=str)),"Saldo":exclusiones.get("Saldo",pd.Series(dtype=float)),"Motivo de exclusión":exclusiones.get("Motivo_exclusion",pd.Series(dtype=str))})
    cq=pd.DataFrame(calidad,columns=["Tipo de inconsistencia","Cantidad","Detalle o referencia de los registros afectados"])
    with pd.ExcelWriter(bio,engine="openpyxl") as w:
        det.to_excel(w,index=False,sheet_name="Asignacion_Diaria"); rg.to_excel(w,index=False,sheet_name="Resumen_Gestores"); rm.to_excel(w,index=False,sheet_name="Resumen_Metas"); ex.to_excel(w,index=False,sheet_name="Exclusiones"); cq.to_excel(w,index=False,sheet_name="Calidad_Datos")
    bio.seek(0); wb=load_workbook(bio)
    header_fill=PatternFill("solid",fgColor="17365D"); header_font=Font(color="FFFFFF",bold=True); warn_fill=PatternFill("solid",fgColor="FCE4D6")
    for ws in wb.worksheets:
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.fill=header_fill; cell.font=header_font; cell.alignment=Alignment(horizontal="center")
        for col in range(1,ws.max_column+1):
            vals=[len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,200)+1)]
            ws.column_dimensions[get_column_letter(col)].width=min(max(vals+[10])+2,45)
        headers={ws.cell(1,c).value:c for c in range(1,ws.max_column+1)}
        for h,c in headers.items():
            hs=str(h).lower()
            if any(k in hs for k in ["monto","saldo","meta","recaudo","brecha","desviacion"]):
                for r in range(2,ws.max_row+1): ws.cell(r,c).number_format='$#,##0;[Red]($#,##0);-'
            if any(k in hs for k in ["porcentaje","cumplimiento"]):
                for r in range(2,ws.max_row+1):
                    # Peso objetivo viene en escala 0-100; convertir formato visual sin alterar el dato no es correcto.
                    ws.cell(r,c).number_format='0.0%' if h not in ["Porcentaje objetivo"] else '0.0'
        if ws.title in ["Exclusiones","Calidad_Datos"]:
            for row in ws.iter_rows(min_row=2):
                for cell in row: cell.fill=warn_fill
    out=BytesIO(); wb.save(out); out.seek(0); return out


def pantalla_configuracion():
    st.header("1. Configuración mensual")
    st.caption("Defina el período, las metas y la capacidad relativa de los gestores.")
    c1,c2=st.columns(2)
    with c1: anio=st.number_input("Año de gestión",min_value=2020,max_value=2100,value=int(st.session_state.anio),step=1)
    with c2: mes=st.selectbox("Mes de gestión",list(MESES),index=list(MESES).index(st.session_state.mes),format_func=lambda x:MESES[x])
    c1,c2=st.columns(2)
    with c1: meta_c=st.number_input("Meta de recaudo - cartera castigada",min_value=0.0,value=float(st.session_state.meta_castigada),step=1000000.0,format="%.0f")
    with c2: meta_v=st.number_input("Meta de recaudo - cartera vencida",min_value=0.0,value=float(st.session_state.meta_vencida),step=1000000.0,format="%.0f")
    st.subheader("Gestores")
    ed=st.data_editor(st.session_state.gestores,num_rows="dynamic",use_container_width=True,key="editor_gestores",
        column_config={"Gestor":st.column_config.TextColumn(required=True),"Prioridad":st.column_config.NumberColumn(min_value=1,step=1),"Peso (%)":st.column_config.NumberColumn(min_value=0.0,max_value=100.0,format="%.2f"),"Activo":st.column_config.CheckboxColumn(default=True)})
    if st.button("Distribuir porcentajes automáticamente"):
        st.session_state.gestores=pesos_automaticos(ed); st.rerun()
    errores=validar_gestores(ed)
    if errores:
        for e in errores: st.warning(e)
    activos=ed[ed.Activo.fillna(False).astype(bool)] if not ed.empty else ed
    st.info(f"Resumen: {MESES[mes]} {anio} · Meta total {formato_pesos(meta_c+meta_v)} · {len(activos)} gestores activos.")
    if st.button("Guardar configuración",type="primary",disabled=bool(errores) or meta_c<=0 or meta_v<=0):
        st.session_state.update({"anio":int(anio),"mes":int(mes),"meta_castigada":float(meta_c),"meta_vencida":float(meta_v),"gestores":ed.copy(),"config_guardada":True})
        st.success("Configuración guardada.")


def pantalla_recaudo():
    st.header("2. Seguimiento de recaudo")
    if not st.session_state.config_guardada:
        st.warning("Complete y guarde primero la configuración mensual."); return
    up=st.file_uploader("Cargue el archivo diario de recaudo (.xlsx)",type=["xlsx","xlsm"],key="recaudo_file")
    if not up: return
    raw=up.getvalue()
    try: sheets=hojas_excel(raw)
    except Exception as e: st.error(f"No fue posible abrir el Excel: {e}"); return
    hoja=st.selectbox("Hoja a analizar",sheets,key="hoja_rec")
    try: df=leer_excel(raw,hoja)
    except Exception as e: st.error(f"No fue posible leer la hoja: {e}"); return
    st.dataframe(df.head(30),use_container_width=True)
    fd=fecha_desde_nombre(up.name)
    corte=st.date_input("Fecha de corte",value=fd or date(st.session_state.anio,st.session_state.mes,1))
    if corte.year!=st.session_state.anio or corte.month!=st.session_state.mes: st.warning("La fecha de corte no corresponde al mes configurado. Puede continuar si el archivo contiene pagos del período configurado.")
    cols=df.columns.tolist()
    with st.form("map_recaudo"):
        st.subheader("Mapeo de columnas")
        m={k:selector_columna(lbl,cols,detectar_columna(cols,k),key=f"rec_{k}") for k,lbl in [("obligacion","Número de obligación"),("cedula","Cédula"),("tipo","Tipo de cartera"),("recaudo","Valor del recaudo"),("fecha_pago","Fecha del pago")]}
        go=st.form_submit_button("Procesar recaudo",type="primary")
    if go:
        with st.spinner("Procesando y validando recaudo..."):
            limpio,cal=procesar_recaudo(df,m,st.session_state.anio,st.session_state.mes)
            res=resumen_recaudo(limpio,st.session_state.meta_castigada,st.session_state.meta_vencida,corte)
            st.session_state.recaudo_limpio=limpio; st.session_state.pagos_obligaciones=set(limpio.Obligacion); st.session_state.pagos_cedulas=set(limpio.Cedula)-{""}; st.session_state.resumen_metas=res; st.session_state.fecha_corte=corte; st.session_state.calidad_recaudo=cal
    if st.session_state.resumen_metas is None: return
    res=st.session_state.resumen_metas; rec=float(res["Recaudo acumulado"].sum()); meta=float(res.Meta.sum()); cump=rec/meta if meta else 0
    ultimo=calendar.monthrange(st.session_state.anio,st.session_state.mes)[1]
    dias_trans=min(st.session_state.fecha_corte.day,ultimo); dias_rest=max(ultimo-dias_trans,0); prom=rec/max(dias_trans,1); requerido=max(meta-rec,0)/max(dias_rest,1) if dias_rest else max(meta-rec,0)
    c=st.columns(5); c[0].metric("Meta consolidada",formato_pesos(meta)); c[1].metric("Recaudo",formato_pesos(rec)); c[2].metric("Cumplimiento",f"{cump:.1%}"); c[3].metric("Promedio diario",formato_pesos(prom)); c[4].metric("Recaudo diario requerido",formato_pesos(requerido))
    st.caption(f"Días transcurridos: {dias_trans} · Días restantes: {dias_rest}")
    show=res.copy(); show["Meta"]=show.Meta.map(formato_pesos); show["Recaudo acumulado"]=show["Recaudo acumulado"].map(formato_pesos); show["Brecha pendiente"]=show["Brecha pendiente"].map(formato_pesos); show["Cumplimiento"]=show.Cumplimiento.map(lambda x:f"{x:.1%}")
    st.dataframe(show,use_container_width=True,hide_index=True)
    sem="🟢 Verde" if cump>=.9 else "🟡 Amarillo" if cump>=.7 else "🔴 Rojo"; st.markdown(f"**Semáforo consolidado:** {sem}")
    a,b=st.columns(2)
    with a: st.plotly_chart(px.bar(res,x="Tipo de cartera",y=["Meta","Recaudo acumulado"],barmode="group",title="Avance frente a la meta"),use_container_width=True)
    with b: st.plotly_chart(px.pie(res,names="Tipo de cartera",values="Recaudo acumulado",title="Recaudo por tipo de cartera"),use_container_width=True)
    bad=[r for r in st.session_state.calidad_recaudo if r[1]>0]
    if bad: st.warning("Se detectaron novedades de calidad. Los registros válidos se procesaron y las novedades quedaron separadas."); st.dataframe(pd.DataFrame(bad,columns=["Inconsistencia","Cantidad","Detalle"]),hide_index=True,use_container_width=True)


def pantalla_asignacion():
    st.header("3. Asignación diaria")
    if st.session_state.resumen_metas is None or st.session_state.recaudo_limpio is None:
        st.warning("Procese primero el recaudo del mes para identificar las cédulas que ya registran pago."); return
    up=st.file_uploader("Cargue el archivo mensual de asignación (.xlsx)",type=["xlsx","xlsm"],key="asig_file")
    if not up: return
    raw=up.getvalue()
    try: sheets=hojas_excel(raw)
    except Exception as e: st.error(f"No fue posible abrir el Excel: {e}"); return
    default=sheets.index("Base") if "Base" in sheets else 0
    hoja=st.selectbox("Hoja de trabajo",sheets,index=default,key="hoja_asig")
    # Algunos libros tienen títulos antes del encabezado; Base del archivo real ya inicia en fila 1.
    df=leer_excel(raw,hoja); st.dataframe(df.head(30),use_container_width=True)
    cols=df.columns.tolist()
    with st.form("map_asig"):
        st.subheader("Mapeo de columnas")
        m={}
        for k,lbl,req in [("obligacion","Número de obligación",True),("cedula","Cédula",True),("cliente","Nombre del cliente",False),("tipo","Tipo de cartera",True),("saldo","Saldo pendiente / capital",True),("gestor","Gestor actual",False)]:
            m[k]=selector_columna(lbl,cols,detectar_columna(cols,k),required=req,key=f"asi_{k}")
        go=st.form_submit_button("Procesar base mensual",type="primary")
    if go:
        with st.spinner("Limpiando base y validando cruces..."):
            base,cal=procesar_asignacion(df,m); st.session_state.asignacion_limpia=base; st.session_state.calidad_asignacion=cal
    if st.session_state.asignacion_limpia is None: return
    admitir=st.checkbox("Admitir obligaciones de cédulas con pagos parciales (quedan con menor prioridad)",value=False)
    e,ex=priorizar_obligaciones(st.session_state.asignacion_limpia,st.session_state.pagos_obligaciones,st.session_state.pagos_cedulas,st.session_state.resumen_metas,admitir)
    st.subheader("Alcance de la asignación")
    modo=st.radio("Seleccione el alcance",["Todas las obligaciones elegibles","Cantidad máxima de obligaciones","Hasta alcanzar un monto total de cartera","Máximo de obligaciones por gestor"],horizontal=True)
    valor=None; max_g=None
    if modo=="Cantidad máxima de obligaciones": valor=st.number_input("Cantidad máxima",min_value=1,value=min(100,max(len(e),1)),step=1)
    elif modo=="Hasta alcanzar un monto total de cartera": valor=st.number_input("Monto total máximo",min_value=1.0,value=float(max(e.Saldo.sum(),1)),step=1000000.0,format="%.0f")
    elif modo=="Máximo de obligaciones por gestor": max_g=st.number_input("Máximo por gestor",min_value=1,value=20,step=1)
    candidatos=aplicar_alcance(e,modo,valor)
    if st.button("Generar asignación diaria",type="primary",disabled=candidatos.empty):
        asig,rg=distribuir_obligaciones(candidatos,st.session_state.gestores,max_g)
        st.session_state.resultado_asignacion=(asig,rg,ex,st.session_state.calidad_recaudo+getattr(st.session_state,"calidad_asignacion",[]))
    if st.session_state.resultado_asignacion is None:
        c=st.columns(4); c[0].metric("Obligaciones elegibles",len(e)); c[1].metric("Saldo elegible",formato_pesos(e.Saldo.sum())); c[2].metric("Excluidas por pago",int((ex.Motivo_exclusion.str.contains("pago",case=False,na=False)).sum())); c[3].metric("Excluidas por inconsistencias",int((ex.Motivo_exclusion.str.contains("Inconsistencia",case=False,na=False)).sum())); return
    asig,rg,ex,cal=st.session_state.resultado_asignacion
    c=st.columns(4); c[0].metric("Obligaciones asignadas",len(asig)); c[1].metric("Saldo asignado",formato_pesos(asig.Saldo.sum())); c[2].metric("Excluidas por pago",int(ex.Motivo_exclusion.str.contains("pago",case=False,na=False).sum())); c[3].metric("Excluidas por inconsistencias",int(ex.Motivo_exclusion.str.contains("Inconsistencia",case=False,na=False).sum()))
    st.caption("La proporcionalidad puede presentar desviaciones porque cada obligación es indivisible. La primera ronda entrega los saldos más altos según la prioridad de gestores.")
    rgshow=rg.copy(); rgshow["Peso (%)"]=rgshow["Peso (%)"].map(lambda x:f"{x:.2f}%"); rgshow["Porcentaje real"]=rgshow["Porcentaje real"].map(lambda x:f"{x:.2%}")
    for col in ["Monto objetivo","Monto asignado","Desviacion","Saldo promedio asignado"]: rgshow[col]=rgshow[col].map(lambda x:formato_pesos(0 if pd.isna(x) else x))
    st.dataframe(rgshow,use_container_width=True,hide_index=True)
    st.plotly_chart(px.bar(rg,x="Gestor",y="Monto asignado",title="Monto asignado por gestor"),use_container_width=True)
    st.subheader("Detalle de asignación"); st.dataframe(asig,use_container_width=True,hide_index=True)
    excel=crear_excel(asig,rg,st.session_state.resumen_metas,ex,cal,date.today())
    nombre=f"Asignacion_Diaria_{date.today():%Y%m%d}.xlsx"
    st.download_button("Descargar asignación diaria en Excel",data=excel,file_name=nombre,mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",type="primary")


init_state()
st.title("📊 Gestión de Recaudo y Asignación Diaria")
st.caption("Aplicación local para seguimiento de metas, priorización de cartera y distribución de cargas de gestión.")
with st.sidebar:
    st.subheader("Estado del proceso")
    st.write("✅ Configuración" if st.session_state.config_guardada else "⬜ Configuración")
    st.write("✅ Recaudo procesado" if st.session_state.recaudo_limpio is not None else "⬜ Recaudo procesado")
    st.write("✅ Base de asignación" if st.session_state.asignacion_limpia is not None else "⬜ Base de asignación")
    if st.button("Reiniciar sesión / nuevo mes"):
        for k in list(st.session_state.keys()): del st.session_state[k]
        st.rerun()

t1,t2,t3=st.tabs(["⚙️ Configuración","💰 Seguimiento de recaudo","📋 Asignación diaria"])
with t1: pantalla_configuracion()
with t2: pantalla_recaudo()
with t3: pantalla_asignacion()
